"""Taxonomy, CSV round-tripping, query filters, URL detection and migration."""

from __future__ import annotations

import datetime
from pathlib import Path

import pytest
from conftest import problem, write_problems

from cpidx import commands, csvio, exporters, importtree
from cpidx.csvio import join_list, split_list
from cpidx.query import Query, QueryError, parse_range
from cpidx.schema import PROBLEM_COLUMNS
from cpidx.taxonomy import Taxonomy, TaxonomyError


# --- taxonomy ----------------------------------------------------------------
def test_descendants_include_the_node_itself(mini_taxonomy: Taxonomy) -> None:
    assert mini_taxonomy.descendants("graphs.traversal") == (
        "graphs.traversal",
        "graphs.traversal.bfs",
        "graphs.traversal.dfs",
    )


def test_ancestors_are_derived_nearest_first(mini_taxonomy: Taxonomy) -> None:
    assert mini_taxonomy.ancestors("graphs.traversal.bfs") == ("graphs.traversal", "graphs")


def test_leaf_detection(mini_taxonomy: Taxonomy) -> None:
    assert mini_taxonomy.is_leaf("graphs.traversal.bfs")
    assert not mini_taxonomy.is_leaf("graphs.traversal")


def test_transitive_prerequisites(mini_taxonomy: Taxonomy) -> None:
    assert mini_taxonomy.transitive_prerequisites("graphs.shortest_path.dijkstra") == {
        "graphs.traversal.bfs"
    }


def test_alias_index_maps_synonyms(mini_taxonomy: Taxonomy) -> None:
    assert mini_taxonomy.alias_index()["breadth first search"] == "graphs.traversal.bfs"


def test_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(TaxonomyError):
        Taxonomy.load(str(tmp_path / "absent.yaml"))


def test_duplicate_ids_rejected(tmp_path: Path) -> None:
    path = tmp_path / "t.yaml"
    path.write_text(
        "version: 0\nnodes:\n- id: a\n  name: A\n  status: canonical\n  description: d\n"
        "- id: a\n  name: A\n  status: canonical\n  description: d\n",
        encoding="utf-8",
    )
    with pytest.raises(TaxonomyError):
        Taxonomy.load(str(path))


# --- the shipped taxonomy ----------------------------------------------------
def test_shipped_taxonomy_is_structurally_clean(real_taxonomy: Taxonomy) -> None:
    assert real_taxonomy.structural_findings() == []


def test_shipped_taxonomy_covers_the_spec_categories(real_taxonomy: Taxonomy) -> None:
    expected = {
        "math",
        "number_theory",
        "combinatorics",
        "geometry",
        "graphs",
        "trees",
        "flows",
        "dp",
        "data_structures",
        "strings",
        "greedy",
        "search",
        "bitwise",
        "game_theory",
        "optimization",
        "randomized",
        "adhoc",
    }
    assert set(real_taxonomy.top_level()) == expected


def test_m0_categories_reach_leaf_depth(real_taxonomy: Taxonomy) -> None:
    """SPEC.md 12 M0: graphs, dp and data_structures to leaf depth."""
    for category in ("graphs", "dp", "data_structures"):
        leaves = [t for t in real_taxonomy.leaves() if t.startswith(category + ".")]
        assert len(leaves) >= 15, category
        assert all(real_taxonomy.get(t).discriminator for t in leaves)


# --- csvio -------------------------------------------------------------------
def test_list_round_trip() -> None:
    assert split_list("a|b|c") == ("a", "b", "c")
    assert split_list("") == ()
    assert join_list(("a", "b")) == "a|b"


def test_render_sorts_by_id_in_byte_order() -> None:
    rows = [problem(id="cses-2"), problem(id="cses-10"), problem(id="cses-1")]
    body = csvio.render(rows, PROBLEM_COLUMNS).splitlines()[1:]
    assert [line.split(",")[0] for line in body] == ["cses-1", "cses-10", "cses-2"]


def test_render_quotes_embedded_commas() -> None:
    text = csvio.render([problem(title="Hello, World")], PROBLEM_COLUMNS)
    assert '"Hello, World"' in text


def test_round_trip_preserves_rows(tmp_path: Path) -> None:
    path = tmp_path / "problems.csv"
    original = [problem(id="cses-1", title="A, B"), problem(id="cses-2", tags="a|b")]
    write_problems(path, original)
    _, rows, findings = csvio.read_problems(str(path))
    assert not findings
    assert [r["title"] for r in rows] == ["A, B", "Example Problem"]


def test_files_use_lf_endings(tmp_path: Path) -> None:
    path = tmp_path / "problems.csv"
    write_problems(path, [problem()])
    assert b"\r\n" not in path.read_bytes()


# --- query -------------------------------------------------------------------
@pytest.mark.parametrize(
    "text,expected",
    [("4", (4, 4)), ("4-6", (4, 6)), ("-6", (1, 6)), ("4-", (4, 10)), ("", None)],
)
def test_parse_range(text: str, expected) -> None:
    assert parse_range(text, 1, 10) == expected


def test_parse_range_rejects_inverted() -> None:
    with pytest.raises(QueryError):
        parse_range("6-4", 1, 10)


def test_tag_filter_matches_descendants(mini_taxonomy: Taxonomy) -> None:
    rows = [
        problem(id="a", primary_tag="graphs.traversal.bfs", tags="graphs.traversal.bfs"),
        problem(id="b", primary_tag="dp.basics.linear", tags="dp.basics.linear"),
    ]
    query = Query(mini_taxonomy)
    query.tag = "graphs"
    assert [r["id"] for r in query.run(rows)] == ["a"]


def test_tag_exact_does_not_match_descendants(mini_taxonomy: Taxonomy) -> None:
    rows = [problem(id="a", primary_tag="graphs.traversal.bfs", tags="graphs.traversal.bfs")]
    query = Query(mini_taxonomy)
    query.tag_exact = "graphs.traversal"
    assert query.run(rows) == []


def test_primary_filter_ignores_secondary_tags(mini_taxonomy: Taxonomy) -> None:
    rows = [
        problem(
            id="a",
            primary_tag="dp.basics.linear",
            tags="dp.basics.linear|graphs.traversal.bfs",
        )
    ]
    query = Query(mini_taxonomy)
    query.primary = "graphs"
    assert query.run(rows) == []
    query.primary = "dp"
    assert len(query.run(rows)) == 1


def test_unknown_tag_raises(mini_taxonomy: Taxonomy) -> None:
    query = Query(mini_taxonomy)
    query.tag = "not.a.tag"
    with pytest.raises(QueryError):
        query.run([])


def test_status_filter_treats_absent_progress_as_untouched(mini_taxonomy: Taxonomy) -> None:
    rows = [problem(id="a"), problem(id="b")]
    progress = [csvio.blank_progress(id="a", status="solved")]
    query = Query(mini_taxonomy)
    query.status = {"untouched"}
    assert [r["id"] for r in query.run(rows, progress)] == ["b"]


def test_missing_prereqs_uses_demonstrated_tags(mini_taxonomy: Taxonomy) -> None:
    rows = [
        problem(id="bfs1", primary_tag="graphs.traversal.bfs", tags="graphs.traversal.bfs"),
        problem(
            id="dij1",
            primary_tag="graphs.shortest_path.dijkstra",
            tags="graphs.shortest_path.dijkstra",
        ),
    ]
    query = Query(mini_taxonomy)
    query.missing_prereqs = True
    # Nothing solved yet, so Dijkstra's BFS prerequisite is unmet.
    assert [r["id"] for r in query.run(rows, [])] == ["dij1"]

    solved = [csvio.blank_progress(id="bfs1", status="solved")]
    query.missing_prereqs = False
    assert {r["id"] for r in query.run(rows, solved)} == {"bfs1", "dij1"}


def test_sort_and_limit(mini_taxonomy: Taxonomy) -> None:
    rows = [problem(id="a", difficulty="7"), problem(id="b", difficulty="2")]
    query = Query(mini_taxonomy)
    assert [r["id"] for r in query.run(rows)] == ["b", "a"]
    query.limit = 1
    assert [r["id"] for r in query.run(rows)] == ["b"]


def test_min_quality_excludes_unrated(mini_taxonomy: Taxonomy) -> None:
    rows = [problem(id="a", quality="5"), problem(id="b", quality="")]
    query = Query(mini_taxonomy)
    query.min_quality = 4
    assert [r["id"] for r in query.run(rows)] == ["a"]


# --- add / detection ---------------------------------------------------------
@pytest.mark.parametrize(
    "url,expected_id,expected_origin",
    [
        ("https://dmoj.ca/problem/ccc21s4", "ccc-2021-s4", "ccc"),
        ("https://dmoj.ca/problem/cco12p1", "cco-2012-p1", "cco"),
        ("https://codeforces.com/contest/1530/problem/F", "cf-1530-f", "codeforces"),
        ("https://cses.fi/problemset/task/1068", "cses-1068", "cses"),
        ("https://oj.uz/problem/view/IOI15_boxes", "ioi-2015-boxes", "ioi"),
        (
            "https://www.usaco.org/index.php?page=viewproblem2&cpid=1070",
            "usaco-cpid1070",
            "usaco",
        ),
    ],
)
def test_url_detection(url: str, expected_id: str, expected_origin: str) -> None:
    detection = commands.detect(url)
    assert detection.fields["id"] == expected_id
    assert detection.fields["origin"] == expected_origin


def test_detection_never_invents_a_tag() -> None:
    detection = commands.detect("https://dmoj.ca/problem/ccc21s4")
    assert "primary_tag" not in detection.fields
    assert "tags" not in detection.fields


def test_build_row_refuses_until_classified() -> None:
    row = commands.build_row("https://dmoj.ca/problem/ccc21s4", "2026-09-19")
    assert "primary_tag" in commands.missing_required(row)


def test_build_row_adds_primary_to_tags() -> None:
    row = commands.build_row(
        "https://dmoj.ca/problem/ccc21s4",
        "2026-09-19",
        {"title": "Daily Commute", "difficulty": "6", "primary_tag": "graphs.traversal.bfs"},
    )
    assert row["tags"] == "graphs.traversal.bfs"
    assert row["difficulty_basis"] == "estimated"
    assert commands.missing_required(row) == []


# --- migration ---------------------------------------------------------------
def test_migration_rewrites_deprecated_tags(mini_taxonomy: Taxonomy) -> None:
    mapping = commands.migration_map(mini_taxonomy)
    assert mapping == {"dp.basics.old_name": "dp.basics.linear"}
    rows = [problem(id="a", primary_tag="dp.basics.old_name", tags="dp.basics.old_name")]
    migrated, changes = commands.migrate_rows(rows, mapping)
    assert migrated[0]["primary_tag"] == "dp.basics.linear"
    assert migrated[0]["tags"] == "dp.basics.linear"
    assert len(changes) == 2


def test_migration_dedupes_when_both_tags_present(mini_taxonomy: Taxonomy) -> None:
    mapping = commands.migration_map(mini_taxonomy)
    rows = [problem(tags="dp.basics.old_name|dp.basics.linear", primary_tag="dp.basics.linear")]
    migrated, _ = commands.migrate_rows(rows, mapping)
    assert migrated[0]["tags"] == "dp.basics.linear"


# --- stats and staleness -----------------------------------------------------
def test_stats_by_origin(mini_taxonomy: Taxonomy) -> None:
    rows = [problem(id="a", origin="cses"), problem(id="b", origin="ccc")]
    assert dict(commands.stats(rows, mini_taxonomy, by="origin")) == {"cses": 1, "ccc": 1}


def test_coverage_gaps_lists_unused_leaves(mini_taxonomy: Taxonomy) -> None:
    gaps = commands.coverage_gaps([problem(tags="graphs.traversal.bfs")], mini_taxonomy)
    assert "graphs.traversal.bfs" not in gaps
    assert "dp.basics.linear" in gaps


def test_stale_rows_respects_max_age() -> None:
    rows = [problem(id="a", verified="2020-01-01"), problem(id="b", verified="2026-09-01")]
    stale = commands.stale_rows(rows, max_age=365, today=datetime.date(2026, 9, 19))
    assert [r["id"] for r, _ in stale] == ["a"]


# --- importer ----------------------------------------------------------------
def test_copy_suffixes_collapse_into_one_candidate(tmp_path: Path) -> None:
    ioi = tmp_path / "IOI" / "25"
    ioi.mkdir(parents=True)
    for name in ("IOI25p1.cpp", "IOI25p1 2.cpp", "IOI25p1 3.cpp"):
        (ioi / name).write_text("int main(){}", encoding="utf-8")
    merged, _ = importtree.derive(str(tmp_path))
    assert list(merged) == ["ioi-2025-p1"]
    assert len(merged["ioi-2025-p1"].sources) == 3
    assert "dedupe" in merged["ioi-2025-p1"].needs


@pytest.mark.parametrize(
    "path,expected_id",
    [
        (("CCC", "S", "21", "21s4.cpp"), "ccc-2021-s4"),
        (("CCC", "J", "10", "2010j1.cpp"), "ccc-2010-j1"),
        (("CCC", "S", "25", "R25s2.cpp"), "ccc-2025-s2"),
        (("USACO", "20", "dec", "P_cowmistry.cpp"), "usaco-2020-dec-platinum-cowmistry"),
        (("CodeForces", "1076", "B_permutation.cpp"), "cf-1076-b"),
        (("CSES", "Tree", "treediameter.cpp"), "cses-treediameter"),
    ],
)
def test_path_patterns(tmp_path: Path, path: tuple[str, ...], expected_id: str) -> None:
    target = tmp_path.joinpath(*path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("int main(){}", encoding="utf-8")
    merged, _ = importtree.derive(str(tmp_path))
    assert expected_id in merged


def test_unlabelled_ccc_file_is_staged_not_dropped(tmp_path: Path) -> None:
    """`dailycommute.cpp` is CCC 2021 S4, but nothing in the path says so."""
    target = tmp_path / "CCC" / "S" / "21" / "dailycommute.cpp"
    target.parent.mkdir(parents=True)
    target.write_text("int main(){}", encoding="utf-8")
    merged, unmatched = importtree.derive(str(tmp_path))
    assert unmatched == []
    candidate = merged["ccc-2021-dailycommute"]
    assert candidate.confidence == "low"
    assert "label" in candidate.needs


def test_reference_directories_are_skipped(tmp_path: Path) -> None:
    target = tmp_path / "Algorithms" / "backtracking" / "minimax.cpp"
    target.parent.mkdir(parents=True)
    target.write_text("int main(){}", encoding="utf-8")
    merged, unmatched = importtree.derive(str(tmp_path))
    assert merged == {} and unmatched == []


def test_staging_never_touches_problems_csv(tmp_path: Path, index) -> None:
    target = Path(index.root) / "CCC" / "S" / "21" / "21s1.cpp"
    target.parent.mkdir(parents=True)
    target.write_text("int main(){}", encoding="utf-8")
    before = Path(index.problems).read_text(encoding="utf-8")
    out = tmp_path / "staging.csv"
    rows, _ = importtree.stage(index.root, str(out), "2026-09-19")
    assert rows
    assert Path(index.problems).read_text(encoding="utf-8") == before


# --- exporters ---------------------------------------------------------------
def test_json_export_derives_ancestors(tmp_path: Path, mini_taxonomy: Taxonomy) -> None:
    import json

    exporters.export_json([problem(tags="graphs.traversal.bfs")], mini_taxonomy, str(tmp_path))
    payload = json.loads((tmp_path / "problems.json").read_text(encoding="utf-8"))
    assert payload["problems"][0]["derived_ancestors"] == ["graphs", "graphs.traversal"]


def test_markdown_export_writes_one_page_per_tag(tmp_path: Path, mini_taxonomy: Taxonomy) -> None:
    exporters.export_markdown([problem(tags="graphs.traversal.bfs")], mini_taxonomy, str(tmp_path))
    page = tmp_path / "by-tag" / "graphs" / "traversal" / "bfs.md"
    assert "Use for unweighted distance." in page.read_text(encoding="utf-8")
    assert (tmp_path / "by-tag" / "index.md").exists()


def test_xlsx_export_has_the_spec_sheets(tmp_path: Path, mini_taxonomy: Taxonomy) -> None:
    from openpyxl import load_workbook

    exporters.export_xlsx([problem(tags="graphs.traversal.bfs")], mini_taxonomy, str(tmp_path))
    book = load_workbook(tmp_path / "problems.xlsx")
    assert "All" in book.sheetnames
    assert "Coverage" in book.sheetnames
    assert "Taxonomy" in book.sheetnames
    assert book["All"].freeze_panes == "A2"


# --- regressions from code review -------------------------------------------
def test_stats_by_difficulty_sorts_numerically(mini_taxonomy: Taxonomy) -> None:
    """Tiers are stored as text; "10" must not sort between "1" and "2"."""
    rows = [
        problem(id="a", difficulty="10"),
        problem(id="b", difficulty="2"),
        problem(id="c", difficulty="1"),
    ]
    assert [k for k, _ in commands.stats(rows, mini_taxonomy, by="difficulty")] == ["1", "2", "10"]


def test_simple_origin_candidates_flag_their_missing_hosts(tmp_path: Path) -> None:
    """`hosts` is required, so a candidate lacking one must not look promotable."""
    target = tmp_path / "COCI" / "10" / "difrencija.cpp"
    target.parent.mkdir(parents=True)
    target.write_text("int main(){}", encoding="utf-8")
    merged, _ = importtree.derive(str(tmp_path))
    candidate = next(iter(merged.values()))
    assert not candidate.fields.get("hosts")
    assert "hosts" in candidate.needs
