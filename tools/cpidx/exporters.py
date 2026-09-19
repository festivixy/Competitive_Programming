"""Generated views of the index.

Normative reference: SPEC.md 10.4 and 11.4. Everything here writes into
`build/`, which is gitignored. These artefacts are never edited and never
committed.
"""

from __future__ import annotations

import json
import os
from collections.abc import Sequence
from typing import TYPE_CHECKING

from . import csvio
from .schema import PROBLEM_COLUMNS

if TYPE_CHECKING:
    from .taxonomy import Taxonomy

TIERS = tuple(range(1, 11))


def _ensure(path: str) -> str:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    return path


def _as_int(value: str | None) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


# --- json --------------------------------------------------------------------
def export_json(rows: Sequence[dict], taxonomy: Taxonomy, out_dir: str) -> list[str]:
    """One document holding the rows with their derived ancestor tags."""
    path = _ensure(os.path.join(out_dir, "problems.json"))
    payload = []
    for row in rows:
        record = {c: row.get(c, "") for c in PROBLEM_COLUMNS}
        for column in ("hosts", "mirror_urls", "tags", "free_tags"):
            record[column] = list(csvio.split_list(row.get(column, "")))
        # Ancestors are derived at export time, never stored: SPEC.md 13.
        ancestors = set()
        for tag in record["tags"]:
            ancestors.update(taxonomy.ancestors(tag))
        record["derived_ancestors"] = sorted(ancestors)
        payload.append(record)
    with open(path, "w", encoding="utf-8", newline="") as handle:
        json.dump({"problems": payload}, handle, indent=2, ensure_ascii=False)
        handle.write("\n")
    return [path]


# --- markdown ----------------------------------------------------------------
def _problem_line(row: dict) -> str:
    bits = ["- "]
    if row.get("url"):
        bits.append("[{}]({})".format(row.get("title") or row["id"], row["url"]))
    else:
        bits.append(row.get("title") or row["id"])
    meta = []
    if row.get("difficulty"):
        meta.append("tier {}".format(row["difficulty"]))
    if row.get("contest"):
        meta.append(row["contest"])
    if row.get("quality"):
        meta.append("quality {}".format(row["quality"]))
    if row.get("canonical") == "true":
        meta.append("**canonical**")
    if meta:
        bits.append(" - " + ", ".join(meta))
    bits.append(" `{}`".format(row["id"]))
    return "".join(bits)


def export_markdown(rows: Sequence[dict], taxonomy: Taxonomy, out_dir: str) -> list[str]:
    """One page per leaf tag under build/by-tag/, per SPEC.md 11.4."""
    by_tag = {}
    for row in rows:
        for tag in csvio.split_list(row.get("tags", "")):
            by_tag.setdefault(tag, []).append(row)

    written = []
    for tag in sorted(by_tag):
        node = taxonomy.get(tag)
        path = _ensure(os.path.join(out_dir, "by-tag", *tag.split(".")) + ".md")
        lines = [f"# {node.name if node else tag}", ""]
        if node:
            if node.description:
                lines += [node.description, ""]
            if node.discriminator:
                lines += ["**When to use.** " + node.discriminator, ""]
            if node.prerequisites:
                lines += [
                    "**Prerequisites:** " + ", ".join(f"`{p}`" for p in node.prerequisites),
                    "",
                ]
        ordered = sorted(by_tag[tag], key=lambda r: (_as_int(r.get("difficulty")) or 0, r["id"]))
        lines.append("## Problems")
        lines.append("")
        lines += [_problem_line(row) for row in ordered]
        lines.append("")
        with open(path, "w", encoding="utf-8", newline="\n") as handle:
            handle.write("\n".join(lines))
        written.append(path)

    index_path = _ensure(os.path.join(out_dir, "by-tag", "index.md"))
    index = ["# Problems by technique", ""]
    for top in taxonomy.top_level():
        tags = [t for t in sorted(by_tag) if taxonomy.top_level_of(t) == top]
        if not tags:
            continue
        node = taxonomy.get(top)
        index.append(f"## {node.name if node else top}")
        index.append("")
        for tag in tags:
            rel = "/".join(tag.split(".")) + ".md"
            node = taxonomy.get(tag)
            index.append(f"- [{node.name if node else tag}]({rel}) - {len(by_tag[tag])} problem(s)")
        index.append("")
    with open(index_path, "w", encoding="utf-8", newline="\n") as handle:
        handle.write("\n".join(index))
    written.append(index_path)
    return written


# --- xlsx --------------------------------------------------------------------
def export_xlsx(rows: Sequence[dict], taxonomy: Taxonomy, out_dir: str) -> list[str]:
    """build/problems.xlsx, per SPEC.md 10.4."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = _ensure(os.path.join(out_dir, "problems.xlsx"))
    workbook = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")

    def add_sheet(title: str, subset: Sequence[dict]):
        sheet = workbook.create_sheet(title[:31])
        sheet.append(list(PROBLEM_COLUMNS))
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in subset:
            sheet.append([row.get(c, "") for c in PROBLEM_COLUMNS])
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(PROBLEM_COLUMNS))
        sheet.auto_filter.ref = f"A1:{last_column}{max(1, len(subset) + 1)}"
        for index, column in enumerate(PROBLEM_COLUMNS, start=1):
            width = max([len(column)] + [len(str(r.get(column, ""))) for r in subset] or [10])
            sheet.column_dimensions[get_column_letter(index)].width = min(48, max(10, width + 2))
        return sheet

    workbook.remove(workbook.active)
    add_sheet("All", rows)

    for top in taxonomy.top_level():
        subset = [
            r
            for r in rows
            if r.get("primary_tag") and taxonomy.top_level_of(r["primary_tag"]) == top
        ]
        if subset:
            add_sheet(taxonomy.get(top).name if taxonomy.get(top) else top, subset)

    # Coverage: leaf tags against difficulty tiers, so gaps are visible.
    coverage = workbook.create_sheet("Coverage")
    coverage.append(["tag"] + [f"tier {t}" for t in TIERS] + ["total"])
    for cell in coverage[1]:
        cell.font = header_font
        cell.fill = header_fill
    counts = {}
    for row in rows:
        tier = _as_int(row.get("difficulty"))
        for tag in csvio.split_list(row.get("tags", "")):
            counts.setdefault(tag, {})
            if tier:
                counts[tag][tier] = counts[tag].get(tier, 0) + 1
    for tag in taxonomy.leaves():
        per_tier = counts.get(tag, {})
        values = [per_tier.get(t, 0) for t in TIERS]
        coverage.append([tag] + values + [sum(values)])
    coverage.freeze_panes = "B2"
    coverage.column_dimensions["A"].width = 52
    last_row = coverage.max_row
    if last_row > 1:
        span = f"B2:{get_column_letter(len(TIERS) + 1)}{last_row}"
        coverage.conditional_formatting.add(
            span,
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="F8696B",
                mid_type="num",
                mid_value=1,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    # Taxonomy: the flattened tree, with the discriminators.
    tax_sheet = workbook.create_sheet("Taxonomy")
    tax_sheet.append(
        ["id", "name", "depth", "leaf", "status", "prerequisites", "description", "discriminator"]
    )
    for cell in tax_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    for tag in sorted(taxonomy.nodes):
        node = taxonomy.get(tag)
        tax_sheet.append(
            [
                node.id,
                node.name,
                node.depth,
                "yes" if taxonomy.is_leaf(tag) else "",
                node.status,
                ", ".join(node.prerequisites),
                node.description,
                node.discriminator,
            ]
        )
    tax_sheet.freeze_panes = "A2"
    for column, width in (("A", 46), ("B", 32), ("F", 40), ("G", 60), ("H", 70)):
        tax_sheet.column_dimensions[column].width = width
    for row in tax_sheet.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(path)
    return [path]


EXPORTERS = {"json": export_json, "md": export_markdown, "xlsx": export_xlsx}
